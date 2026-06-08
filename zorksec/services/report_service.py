"""Reporting engine.

Builds SOC reports from a simple, structured document model and renders them to
multiple formats. Markdown, HTML, JSON, and CSV are always available (pure
Python). DOCX and PDF are best-effort: if the optional library is missing the
engine falls back gracefully and tells the caller which formats are available.

Report types: incident, threat_hunt, dfir, assessment, executive.
"""

from __future__ import annotations

import csv
import datetime as _dt
import io
import json
from dataclasses import asdict, dataclass, field

from sqlalchemy.orm import Session

from zorksec.config import Settings, ensure_directories, get_settings
from zorksec.repositories.report_repository import ReportRepository
from zorksec.utils.logging import get_logger

logger = get_logger(__name__)

REPORT_TYPES = ["incident", "threat_hunt", "dfir", "assessment", "executive"]
TEXT_FORMATS = ["markdown", "html", "json", "csv"]


@dataclass
class ReportSection:
    heading: str
    body: str = ""
    bullets: list[str] = field(default_factory=list)


@dataclass
class ReportDocument:
    title: str
    report_type: str
    author: str = "ZorkSec"
    summary: str = ""
    sections: list[ReportSection] = field(default_factory=list)
    created_at: str = field(default_factory=lambda: _dt.datetime.utcnow().isoformat(timespec="seconds"))

    def to_dict(self) -> dict:
        return {
            "title": self.title,
            "report_type": self.report_type,
            "author": self.author,
            "summary": self.summary,
            "created_at": self.created_at,
            "sections": [asdict(s) for s in self.sections],
        }


def _esc(text: str) -> str:
    return (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def render_markdown(doc: ReportDocument) -> str:
    lines = [f"# {doc.title}", "",
             f"*Type:* {doc.report_type}  ", f"*Author:* {doc.author}  ",
             f"*Generated:* {doc.created_at}", ""]
    if doc.summary:
        lines += ["## Summary", "", doc.summary, ""]
    for section in doc.sections:
        lines.append(f"## {section.heading}")
        lines.append("")
        if section.body:
            lines += [section.body, ""]
        for bullet in section.bullets:
            lines.append(f"- {bullet}")
        if section.bullets:
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def render_html(doc: ReportDocument) -> str:
    parts = [
        "<!DOCTYPE html><html><head><meta charset='utf-8'>",
        f"<title>{_esc(doc.title)}</title>",
        "<style>body{font-family:system-ui,Segoe UI,sans-serif;max-width:820px;"
        "margin:40px auto;color:#1c2230;line-height:1.5}h1{color:#3b2a7a}"
        "h2{border-bottom:1px solid #ddd;padding-bottom:4px}"
        ".meta{color:#666;font-size:13px}</style></head><body>",
        f"<h1>{_esc(doc.title)}</h1>",
        f"<p class='meta'>Type: {_esc(doc.report_type)} &middot; Author: "
        f"{_esc(doc.author)} &middot; Generated: {_esc(doc.created_at)}</p>",
    ]
    if doc.summary:
        parts.append(f"<h2>Summary</h2><p>{_esc(doc.summary)}</p>")
    for section in doc.sections:
        parts.append(f"<h2>{_esc(section.heading)}</h2>")
        if section.body:
            parts.append(f"<p>{_esc(section.body)}</p>")
        if section.bullets:
            parts.append("<ul>" + "".join(f"<li>{_esc(b)}</li>" for b in section.bullets) + "</ul>")
    parts.append("</body></html>")
    return "".join(parts)


def render_json(doc: ReportDocument) -> str:
    return json.dumps(doc.to_dict(), indent=2)


def render_csv(doc: ReportDocument) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["section", "type", "content"])
    if doc.summary:
        writer.writerow(["Summary", "body", doc.summary])
    for section in doc.sections:
        if section.body:
            writer.writerow([section.heading, "body", section.body])
        for bullet in section.bullets:
            writer.writerow([section.heading, "bullet", bullet])
    return buf.getvalue()


def available_formats() -> list[str]:
    """Return all renderable formats, including optional ones if installed."""
    formats = list(TEXT_FORMATS)
    try:
        import docx  # noqa: F401  (python-docx)
        formats.append("docx")
    except ImportError:
        pass
    try:
        import reportlab  # noqa: F401
        formats.append("pdf")
    except ImportError:
        pass
    return formats


def render(doc: ReportDocument, fmt: str) -> tuple[bytes, str]:
    """Render to (content_bytes, file_extension). Raises ValueError if unsupported."""
    fmt = fmt.lower()
    if fmt == "markdown":
        return render_markdown(doc).encode("utf-8"), "md"
    if fmt == "html":
        return render_html(doc).encode("utf-8"), "html"
    if fmt == "json":
        return render_json(doc).encode("utf-8"), "json"
    if fmt == "csv":
        return render_csv(doc).encode("utf-8"), "csv"
    if fmt == "docx":
        return _render_docx(doc), "docx"
    if fmt == "pdf":
        return _render_pdf(doc), "pdf"
    raise ValueError(f"Unsupported format '{fmt}'. Available: {available_formats()}")


def _render_docx(doc: ReportDocument) -> bytes:
    try:
        from docx import Document  # type: ignore
    except ImportError as exc:
        raise ValueError("DOCX export requires 'python-docx'. Install it or use markdown/html.") from exc
    document = Document()
    document.add_heading(doc.title, level=0)
    document.add_paragraph(f"Type: {doc.report_type} | Author: {doc.author} | {doc.created_at}")
    if doc.summary:
        document.add_heading("Summary", level=1)
        document.add_paragraph(doc.summary)
    for section in doc.sections:
        document.add_heading(section.heading, level=1)
        if section.body:
            document.add_paragraph(section.body)
        for bullet in section.bullets:
            document.add_paragraph(bullet, style="List Bullet")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _render_pdf(doc: ReportDocument) -> bytes:
    try:
        from reportlab.lib.pagesizes import LETTER  # type: ignore
        from reportlab.pdfgen import canvas  # type: ignore
    except ImportError as exc:
        raise ValueError("PDF export requires 'reportlab'. Install it or use html/markdown.") from exc
    buf = io.BytesIO()
    pdf = canvas.Canvas(buf, pagesize=LETTER)
    width, height = LETTER
    y = height - 50
    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(50, y, doc.title)
    y -= 24
    pdf.setFont("Helvetica", 9)
    pdf.drawString(50, y, f"Type: {doc.report_type} | Author: {doc.author} | {doc.created_at}")
    y -= 24

    def line(text: str, size: int = 11, bold: bool = False) -> None:
        nonlocal y
        if y < 60:
            pdf.showPage()
            y = height - 50
        pdf.setFont("Helvetica-Bold" if bold else "Helvetica", size)
        pdf.drawString(50, y, text[:110])
        y -= size + 6

    if doc.summary:
        line("Summary", 13, bold=True)
        line(doc.summary)
    for section in doc.sections:
        line(section.heading, 13, bold=True)
        if section.body:
            line(section.body)
        for bullet in section.bullets:
            line(f"- {bullet}")
    pdf.showPage()
    pdf.save()
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Templates: build a structured ReportDocument for each report type.
# ---------------------------------------------------------------------------
def build_template(report_type: str, title: str, context: dict | None = None) -> ReportDocument:
    context = context or {}
    report_type = report_type.lower()
    if report_type not in REPORT_TYPES:
        raise ValueError(f"Unknown report type '{report_type}'. Choose from {REPORT_TYPES}.")

    summary = context.get("summary", "")
    if report_type == "incident":
        sections = [
            ReportSection("Timeline", context.get("timeline", "Describe events in order.")),
            ReportSection("Impact", context.get("impact", "What was affected?")),
            ReportSection("Indicators of Compromise", bullets=context.get("iocs", [])),
            ReportSection("Containment & Eradication", context.get("response", "")),
            ReportSection("Recommendations", bullets=context.get("recommendations", [])),
        ]
    elif report_type == "threat_hunt":
        sections = [
            ReportSection("Hypothesis", context.get("hypothesis", "What were you looking for?")),
            ReportSection("Data Sources", bullets=context.get("data_sources", [])),
            ReportSection("Findings", context.get("findings", "")),
            ReportSection("ATT&CK Techniques", bullets=context.get("attack", [])),
        ]
    elif report_type == "dfir":
        sections = [
            ReportSection("Evidence Acquired", bullets=context.get("evidence", [])),
            ReportSection("Analysis", context.get("analysis", "")),
            ReportSection("Conclusions", context.get("conclusions", "")),
        ]
    elif report_type == "assessment":
        sections = [
            ReportSection("Scope", context.get("scope", "")),
            ReportSection("Findings", bullets=context.get("findings", [])),
            ReportSection("Risk Rating", context.get("risk", "")),
            ReportSection("Remediation", bullets=context.get("remediation", [])),
        ]
    else:  # executive
        sections = [
            ReportSection("Business Impact", context.get("impact", "")),
            ReportSection("Key Risks", bullets=context.get("risks", [])),
            ReportSection("Next Steps", bullets=context.get("next_steps", [])),
        ]
    return ReportDocument(
        title=title,
        report_type=report_type,
        author=context.get("author", "ZorkSec"),
        summary=summary,
        sections=sections,
    )


class ReportService:
    def __init__(self, session: Session, settings: Settings | None = None) -> None:
        self.session = session
        self.settings = settings or get_settings()
        self.reports = ReportRepository(session)

    @staticmethod
    def report_types() -> list[str]:
        return list(REPORT_TYPES)

    @staticmethod
    def formats() -> list[str]:
        return available_formats()

    def generate(self, report_type: str, title: str, context: dict | None = None,
                 fmt: str = "markdown") -> Report:
        """Build, render, persist (store the text body), and return the Report."""
        doc = build_template(report_type, title, context)
        # Persist a markdown body for portability regardless of export format.
        body = render_markdown(doc)
        report = self.reports.create(title=title, report_type=report_type, body=body, fmt=fmt)
        logger.info("Generated %s report '%s' (id=%s)", report_type, title, report.id)
        return report

    def export_to_file(self, report_type: str, title: str, context: dict | None,
                       fmt: str) -> str:
        """Render and write a report to the reports directory; return the path."""
        doc = build_template(report_type, title, context)
        content, ext = render(doc, fmt)
        ensure_directories(self.settings)
        safe = "".join(c if c.isalnum() or c in "-_" else "_" for c in title)[:60]
        stamp = _dt.datetime.utcnow().strftime("%Y%m%d-%H%M%S")
        path = self.settings.reports_dir / f"{safe}-{stamp}.{ext}"
        path.write_bytes(content)
        return str(path)
